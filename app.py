import os
import io
import glob
import sqlite3
import base64
from datetime import date, datetime
import pandas as pd
from PIL import Image
import streamlit as st

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(
    page_title="Suivi Chantier & Étanchéité",
    page_icon="🏗️",
    layout="centered",
    initial_sidebar_state="collapsed",
)

st.markdown("""
<style>
    @media (max-width: 768px) {
        .block-container {
            padding-top: 1rem !important;
            padding-bottom: 2rem !important;
            padding-left: 0.6rem !important;
            padding-right: 0.6rem !important;
        }
        h1 { font-size: 1.4rem !important; }
        h2 { font-size: 1.2rem !important; }
        h3 { font-size: 1.1rem !important; }
        .stButton>button {
            height: 48px !important;
            font-size: 15px !important;
            font-weight: 600 !important;
            border-radius: 10px !important;
            margin-top: 4px !important;
            margin-bottom: 4px !important;
        }
    }
    
    .worker-frame {
        background-color: #111827;
        border: 1.5px solid #374151;
        border-radius: 16px;
        padding: 16px 14px;
        margin-bottom: 22px;
        box-shadow: 0 4px 10px rgba(0, 0, 0, 0.35);
        transition: all 0.2s ease-in-out;
    }
    
    .worker-frame:hover {
        border-color: #3B82F6;
    }

    .worker-header {
        display: flex;
        align-items: center;
        gap: 12px;
        margin-bottom: 12px;
        padding-bottom: 10px;
        border-bottom: 1px solid #1F2937;
    }

    .avatar-img {
        width: 65px !important;
        height: 65px !important;
        border-radius: 50% !important;
        object-fit: cover !important;
        border: 2px solid #3B82F6 !important;
        box-shadow: 0 2px 6px rgba(0, 0, 0, 0.3) !important;
        display: block;
        margin: auto;
    }

    .avatar-placeholder {
        width: 65px;
        height: 65px;
        border-radius: 50%;
        background: #1F2937;
        border: 2px solid #64748B;
        display: flex;
        align-items: center;
        justify-content: center;
        font-size: 30px;
        margin: auto;
    }
</style>
""", unsafe_allow_html=True)

DB_PATH = "chantier_master.db"
PHOTOS_DIR = "photos"
os.makedirs(PHOTOS_DIR, exist_ok=True)

ADMIN_PASSWORD = "admin"

LISTE_CHANTIERS_INIT = [
    "CAC-31-24", "CMA-09-23", "ECOLE SBA", "GROUPEMENT GENDARMERIE",
    "HOB-342-08-24", "HTA ZONE SBA", "IZM-31-24(CES)", "LBU-31-23",
    "LCP-31-24", "MARAVAL", "VILLA CITE EL RIAD ORAN",
    "VILLA Hasnaoui MAKAM", "VILLA Hasnaoui Outhman", "ESC-16-24",
    "EN ATTENTE / DEPOT"
]

LISTE_TACHES_INIT = [
    "PAX", "PARE-VAPEUR", "SOKLE PARE-VAPEUR", "ELASTOTEK",
    "ELASTOTEK GOURGE", "ELASTOTEK JOINTAGE", "ELASTOTEK LEVI",
    "ELASTOTEK NETTOYAGE", "Forme de pente", "GOURGE",
    "JOINT DE DILATATION", "MASTIC", "nettoyage", "PONSAGE",
    "décapage", "BACHE A EAU", "PISCINE", "SOUS CARRELAGE",
    "TEST EAU", "coupe-feu", "Couvre-joint", "DALLE Cheminée",
    "regard", "Traitement de l'ascenseur", "BRICOL", "BRICOL ELASTOTEK",
    "BRICOL SILICONE", "BRICOL SOKLE", "BRICOL SOUS CARRELAGE",
    "BRICOL PARE-VAPEUR", "DIVERS"
]

EFFECTIF_GLOBAL_INIT = [
    "ADDA Abbess", "MEKHACHEF DJAMEL", "MESTEFAOUI AHMED", "ARGOUB HALIM",
    "FEHIM CHIBANI AZZOUZ", "TAIBI REDA", "ABED OMAR", "ZEGHDAN ABDELKADER",
    "BAGHDADI ALI", "BOUKHELIF KAMEL", "MOKHTARI Omar", "GHEZINI Habib",
    "BERACHEMI AHMED", "ARAR AISSA", "MOKHTARI Djelloul", "HAFDI Rachid",
    "BENHAMMADI Mohamed", "MOUISSI WALID", "BENSEMICHA MEROUANE",
    "TOUATI Zouaoui", "GHRIBI MOHAMED", "MESSAOUDI ABDELKRIM",
    "TAHAR BOUZIAN YOUCEF"
]

def get_db_connection():
    return sqlite3.connect(DB_PATH, timeout=15)

def init_database():
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS chantiers_ref (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nom TEXT NOT NULL UNIQUE
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS taches_ref (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nom TEXT NOT NULL UNIQUE
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS workers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nom TEXT NOT NULL UNIQUE,
            chantier_fixe TEXT NOT NULL DEFAULT 'CAC-31-24'
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS conducteurs_meta (
            tag TEXT PRIMARY KEY,
            nom_affiche TEXT NOT NULL
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS conducteur_chantiers (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            conducteur_tag TEXT NOT NULL,
            chantier TEXT NOT NULL,
            UNIQUE(conducteur_tag, chantier)
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS pointages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            date_jour TEXT NOT NULL,
            chantier TEXT NOT NULL,
            conducteur TEXT NOT NULL DEFAULT 'Conducteur',
            worker_id INTEGER NOT NULL,
            statut TEXT,
            tache TEXT,
            quantite REAL,
            unite TEXT,
            appreciation TEXT,
            observation TEXT,
            score REAL,
            UNIQUE(date_jour, chantier, worker_id)
        )
    """)

    for ch in LISTE_CHANTIERS_INIT:
        c.execute("INSERT OR IGNORE INTO chantiers_ref (nom) VALUES (?)", (ch,))
    for tch in LISTE_TACHES_INIT:
        c.execute("INSERT OR IGNORE INTO taches_ref (nom) VALUES (?)", (tch,))
    for w in EFFECTIF_GLOBAL_INIT:
        c.execute("INSERT OR IGNORE INTO workers (nom, chantier_fixe) VALUES (?, ?)", (w, 'CAC-31-24'))

    c.execute("INSERT OR IGNORE INTO conducteurs_meta (tag, nom_affiche) VALUES ('c1', 'Conducteur 1')")
    c.execute("INSERT OR IGNORE INTO conducteurs_meta (tag, nom_affiche) VALUES ('c2', 'Conducteur 2')")
    
    conn.commit()
    conn.close()

init_database()

def get_all_chantiers():
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT nom FROM chantiers_ref ORDER BY nom ASC")
    res = [row[0] for row in c.fetchall()]
    conn.close()
    return res

def get_all_taches():
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT nom FROM taches_ref ORDER BY nom ASC")
    res = [row[0] for row in c.fetchall()]
    conn.close()
    return res

def get_conducteurs_dict():
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT tag, nom_affiche FROM conducteurs_meta")
    res = {row[0]: row[1] for row in c.fetchall()}
    conn.close()
    return res

def get_workers_df():
    conn = get_db_connection()
    df = pd.read_sql_query("SELECT id, nom, chantier_fixe FROM workers ORDER BY nom ASC", conn)
    conn.close()
    return df

def get_chantiers_conducteur(conducteur_tag):
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT chantier FROM conducteur_chantiers WHERE conducteur_tag = ? ORDER BY chantier ASC", (conducteur_tag,))
    lignes = c.fetchall()
    conn.close()
    return [r[0] for r in lignes]

def est_deja_valide(date_str, chantier):
    conn = get_db_connection()
    c = conn.cursor()
    c.execute("SELECT COUNT(*) FROM pointages WHERE date_jour = ? AND chantier = ?", (date_str, chantier))
    count = c.fetchone()[0]
    conn.close()
    return count > 0

def get_photo_path(identifiant):
    nom_clean = identifiant.replace(" ", "_")
    for ext in [".jpg", ".jpeg", ".png"]:
        p = os.path.join(PHOTOS_DIR, f"{nom_clean}{ext}")
        if os.path.exists(p):
            return p
    return None

def get_avatar_html(identifiant):
    path = get_photo_path(identifiant)
    if path and os.path.exists(path):
        with open(path, "rb") as f:
            data = base64.b64encode(f.read()).decode("utf-8")
        ext = path.split(".")[-1].lower()
        mime = "jpeg" if ext in ["jpg", "jpeg"] else "png"
        return f'<img src="data:image/{mime};base64,{data}" class="avatar-img" />'
    return '<div class="avatar-placeholder">👷</div>'

def generer_classeur_pro_excel(df_data, titre_rapport):
    wb = openpyxl.Workbook()
    
    BLEU_HEADER = "1E3A8A"
    BLEU_TITRE = "0F172A"
    VERT_TITRE = "065F46"
    VERT_HEADER = "047857"
    GRIS_ZEBRA = "F8FAFC"
    GRIS_TOTAL = "E2E8F0"
    BORDER_COLOR = "CBD5E1"
    
    VERT_BG, VERT_TXT = "DCFCE7", "166534"
    JAUNE_BG, JAUNE_TXT = "FEF9C3", "854D0E"
    ORANGE_BG, ORANGE_TXT = "FFEDD5", "C2410C"
    ROUGE_BG, ROUGE_TXT = "FEE2E2", "991B1B"
    
    font_titre = Font(name="Segoe UI", size=13, bold=True, color="FFFFFF")
    font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_body = Font(name="Segoe UI", size=9)
    font_total = Font(name="Segoe UI", size=10, bold=True)
    
    fill_header = PatternFill(start_color=BLEU_HEADER, end_color=BLEU_HEADER, fill_type="solid")
    fill_titre = PatternFill(start_color=BLEU_TITRE, end_color=BLEU_TITRE, fill_type="solid")
    fill_zebra = PatternFill(start_color=GRIS_ZEBRA, end_color=GRIS_ZEBRA, fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fill_total = PatternFill(start_color=GRIS_TOTAL, end_color=GRIS_TOTAL, fill_type="solid")
    
    thin_side = Side(style='thin', color=BORDER_COLOR)
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    
    ws_sum = wb.active
    ws_sum.title = "Synthèse Chantiers"
    ws_sum.views.sheetView[0].showGridLines = True
    
    ws_sum.merge_cells("A1:G1")
    ws_sum["A1"] = f"RÉCAPITULATIF DES CHANTIERS — {titre_rapport.upper()}"
    ws_sum["A1"].font = font_titre
    ws_sum["A1"].fill = fill_titre
    ws_sum["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_sum.row_dimensions[1].height = 40
    
    headers_sum = ["Chantier", "Effectif Actif", "Jours Validés", "Production (m²)", "Total Lignes", "Dernier Pointage", "Conducteur"]
    ws_sum.append([])
    ws_sum.append(headers_sum)
    ws_sum.row_dimensions[3].height = 28
    
    for col_idx in range(1, len(headers_sum) + 1):
        c = ws_sum.cell(row=3, column=col_idx)
        c.font = font_header
        c.fill = fill_header
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = cell_border
        
    row_idx = 4
    tot_ouv = 0
    tot_j = 0.0
    tot_m2 = 0.0
    
    for ch_name, grp in df_data.groupby("Chantier"):
        ouv_c = grp["Ouvrier"].nunique()
        p_c = sum(1 for v in grp["Statut"] if "Présent" in str(v))
        d_c = sum(1 for v in grp["Statut"] if "1/2" in str(v))
        j_c = p_c + (d_c * 0.5)
        m2_c = grp[(grp["Unite"] == "m²") & (grp["Quantite"] > 0)]["Quantite"].sum()
        last_date = grp["Date"].max()
        cond_nom = grp["Conducteur"].iloc[-1]
        
        ws_sum.append([ch_name, ouv_c, j_c, round(m2_c, 1), len(grp), last_date, cond_nom])
        ws_sum.row_dimensions[row_idx].height = 22
        
        cur_fill = fill_zebra if row_idx % 2 == 0 else fill_white
        for col_idx in range(1, len(headers_sum) + 1):
            c = ws_sum.cell(row=row_idx, column=col_idx)
            c.font = font_body
            c.border = cell_border
            c.fill = cur_fill
            if col_idx in [2, 3, 4, 5, 6]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
                
        tot_ouv += ouv_c
        tot_j += j_c
        tot_m2 += m2_c
        row_idx += 1
        
    ws_sum.append(["TOTAL GÉNÉRAL", tot_ouv, tot_j, round(tot_m2, 1), len(df_data), "-", "-"])
    ws_sum.row_dimensions[row_idx].height = 26
    for col_idx in range(1, len(headers_sum) + 1):
        c = ws_sum.cell(row=row_idx, column=col_idx)
        c.font = font_total
        c.fill = fill_total
        c.border = cell_border
        if col_idx in [2, 3, 4, 5]:
            c.alignment = Alignment(horizontal="center", vertical="center")

    largeurs_sum = {"A": 24, "B": 16, "C": 16, "D": 18, "E": 15, "F": 18, "G": 20}
    for col_lettre, larg in largeurs_sum.items():
        ws_sum.column_dimensions[col_lettre].width = larg

    ws_ouv = wb.create_sheet(title="Bilan & Primes Ouvriers")
    ws_ouv.views.sheetView[0].showGridLines = True
    
    ws_ouv.merge_cells("A1:J1")
    ws_ouv["A1"] = f"BILAN INDIVIDUEL MENSUEL & ÉVALUATION DES PRIMES — {titre_rapport.upper()}"
    ws_ouv["A1"].font = font_titre
    ws_ouv["A1"].fill = PatternFill(start_color=VERT_TITRE, end_color=VERT_TITRE, fill_type="solid")
    ws_ouv["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_ouv.row_dimensions[1].height = 40
    
    headers_ouvriers = [
        "Nom et Prénom", "Chantier Principal", "Jours Payés", "Absences Autorisées",
        "Absences Injustifiées", "Production (m²)", "Score Moyen", "Éligibilité Prime",
        "Proposition Décision", "Montant Prime (DZD)"
    ]
    ws_ouv.append([])
    ws_ouv.append(headers_ouvriers)
    ws_ouv.row_dimensions[3].height = 28
    
    for col_idx in range(1, len(headers_ouvriers) + 1):
        c = ws_ouv.cell(row=3, column=col_idx)
        c.font = font_header
        c.fill = PatternFill(start_color=VERT_HEADER, end_color=VERT_HEADER, fill_type="solid")
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = cell_border

    largeurs_ouv = {
        "A": 26, "B": 20, "C": 14, "D": 18,
        "E": 20, "F": 16, "G": 14, "H": 22,
        "I": 24, "J": 22
    }
    for col_lettre, larg in largeurs_ouv.items():
        ws_ouv.column_dimensions[col_lettre].width = larg

    r_ouv_idx = 4
    tot_j_ouv = 0.0
    tot_m2_ouv = 0.0

    for w_name, w_grp in df_data.groupby("Ouvrier"):
        ch_princip = w_grp["Chantier"].mode()[0] if not w_grp.empty else "-"
        p_cnt = sum(1 for v in w_grp["Statut"] if "Présent" in str(v))
        d_cnt = sum(1 for v in w_grp["Statut"] if "1/2" in str(v))
        j_payes = p_cnt + (d_cnt * 0.5)
        
        abs_aut = sum(1 for v in w_grp["Statut"] if "Autorisée" in str(v))
        abs_injust = sum(1 for v in w_grp["Statut"] if "Non Autorisée" in str(v) or "Injustifiée" in str(v))
        
        prod_m2 = w_grp[(w_grp["Unite"] == "m²") & (w_grp["Quantite"] > 0)]["Quantite"].sum()
        
        scores_valides = pd.to_numeric(w_grp["Score"], errors='coerce').dropna()
        score_moy = scores_valides.mean() if not scores_valides.empty else 0.0
        
        if abs_injust > 0:
            elig_prime = "Non Éligible"
            prop_decision = "Pénalité (Absence injustifiée)"
        elif score_moy >= 90.0 and j_payes >= 20.0:
            elig_prime = "Éligible Prime Maximale (A+)"
            prop_decision = "Prime de Rendement Supérieure"
        elif score_moy >= 80.0:
            elig_prime = "Éligible Prime Standard (A)"
            prop_decision = "Prime de Rendement Normale"
        elif score_moy >= 70.0:
            elig_prime = "Prime d'Encouragement (B)"
            prop_decision = "Prime Partielle / Encouragement"
        else:
            elig_prime = "Sans Prime"
            prop_decision = "Rendement Insuffisant"

        vals_ouvrier = [
            w_name, ch_princip, j_payes, abs_aut, abs_injust,
            round(prod_m2, 1), round(score_moy, 1), elig_prime, prop_decision, ""
        ]
        ws_ouv.append(vals_ouvrier)
        ws_ouv.row_dimensions[r_ouv_idx].height = 22
        
        c_fill_def = fill_zebra if r_ouv_idx % 2 == 0 else fill_white
        for col_idx in range(1, len(headers_ouvriers) + 1):
            c = ws_ouv.cell(row=r_ouv_idx, column=col_idx)
            c.font = font_body
            c.border = cell_border
            c.fill = c_fill_def
            
            if col_idx in [3, 4, 5, 6, 7]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [8, 9]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 10:
                c.alignment = Alignment(horizontal="right", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
                
            if col_idx == 8:
                if "A+" in elig_prime or "Maximale" in elig_prime:
                    c.fill = PatternFill(start_color=VERT_BG, end_color=VERT_BG, fill_type="solid")
                    c.font = Font(name="Segoe UI", size=9, bold=True, color=VERT_TXT)
                elif "Standard" in elig_prime:
                    c.fill = PatternFill(start_color=JAUNE_BG, end_color=JAUNE_BG, fill_type="solid")
                    c.font = Font(name="Segoe UI", size=9, bold=True, color=JAUNE_TXT)
                elif "Encouragement" in elig_prime:
                    c.fill = PatternFill(start_color=ORANGE_BG, end_color=ORANGE_BG, fill_type="solid")
                    c.font = Font(name="Segoe UI", size=9, bold=True, color=ORANGE_TXT)
                else:
                    c.fill = PatternFill(start_color=ROUGE_BG, end_color=ROUGE_BG, fill_type="solid")
                    c.font = Font(name="Segoe UI", size=9, bold=True, color=ROUGE_TXT)

        tot_j_ouv += j_payes
        tot_m2_ouv += prod_m2
        r_ouv_idx += 1

    ws_ouv.append(["TOTAL / MOYENNE", f"{df_data['Ouvrier'].nunique()} Ouvriers", tot_j_ouv, "-", "-", round(tot_m2_ouv, 1), "-", "-", "-", ""])
    ws_ouv.row_dimensions[r_ouv_idx].height = 26
    for col_idx in range(1, len(headers_ouvriers) + 1):
        c = ws_ouv.cell(row=r_ouv_idx, column=col_idx)
        c.font = font_total
        c.fill = fill_total
        c.border = cell_border
        if col_idx in [3, 4, 5, 6, 7]:
            c.alignment = Alignment(horizontal="center", vertical="center")

    headers_detail = ["Date", "Conducteur", "Ouvrier", "Statut", "Corps d'état / Tâche", "Production", "Contrôle Qualité", "Observation / Rendement", "Score"]
    largeurs_detail = {
        "A": 14, "B": 18, "C": 26, "D": 20, "E": 28,
        "F": 16, "G": 22, "H": 38, "I": 12
    }

    for ch_name, grp in df_data.groupby("Chantier"):
        safe_title = ch_name.replace("/", "-").replace("\\", "-")[:28]
        ws_ch = wb.create_sheet(title=safe_title)
        ws_ch.views.sheetView[0].showGridLines = True
        
        ws_ch.merge_cells("A1:I1")
        ws_ch["A1"] = f"CHANTIER : {ch_name.upper()} — RAPPORT D'ACTIVITÉ"
        ws_ch["A1"].font = font_titre
        ws_ch["A1"].fill = fill_titre
        ws_ch["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws_ch.row_dimensions[1].height = 38
        
        ouv_c = grp["Ouvrier"].nunique()
        p_c = sum(1 for v in grp["Statut"] if "Présent" in str(v))
        d_c = sum(1 for v in grp["Statut"] if "1/2" in str(v))
        j_c = p_c + (d_c * 0.5)
        m2_c = grp[(grp["Unite"] == "m²") & (grp["Quantite"] > 0)]["Quantite"].sum()
        
        ws_ch.merge_cells("A2:D2")
        ws_ch["A2"] = f"Effectif : {ouv_c} ouvrier(s) | Jours Payés : {j_c} j"
        ws_ch["A2"].font = Font(name="Segoe UI", size=10, bold=True, color="1E3A8A")
        ws_ch["A2"].alignment = Alignment(horizontal="left", vertical="center")
        
        ws_ch.merge_cells("E2:I2")
        ws_ch["E2"] = f"Production Métrée : {round(m2_c, 1)} m² | Total Saisies : {len(grp)}"
        ws_ch["E2"].font = Font(name="Segoe UI", size=10, bold=True, color="475569")
        ws_ch["E2"].alignment = Alignment(horizontal="right", vertical="center")
        ws_ch.row_dimensions[2].height = 24
        
        ws_ch.append([])
        ws_ch.append(headers_detail)
        ws_ch.row_dimensions[4].height = 28
        
        for col_idx in range(1, len(headers_detail) + 1):
            c = ws_ch.cell(row=4, column=col_idx)
            c.font = font_header
            c.fill = fill_header
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = cell_border
            
        r_idx = 5
        for _, r in grp.iterrows():
            prod_aff = "Bricolage" if r["Unite"] == "Sans métrage" else (f"{r['Quantite']} {r['Unite']}" if r["Quantite"] > 0 else "-")
            statut_aff = str(r["Statut"]).replace("(Journée)", "").replace("(Congé/Maladie)", "").replace("🟢", "").replace("🟡", "").replace("🔴", "").strip()
            qual_clean = str(r["Qualité"]).replace("/ À reprendre", "").replace("🟢", "").replace("🟡", "").replace("🔴", "").strip()
            obs_clean = str(r["Observation"]).replace("🟢", "").replace("🟡", "").replace("🔴", "").strip()
            
            vals = [
                r["Date"], r["Conducteur"], r["Ouvrier"], statut_aff,
                r["Tâche"], prod_aff, qual_clean, obs_clean, r["Score"]
            ]
            ws_ch.append(vals)
            ws_ch.row_dimensions[r_idx].height = 22
            
            c_fill_defaut = fill_zebra if r_idx % 2 == 0 else fill_white
            for col_idx in range(1, len(headers_detail) + 1):
                c = ws_ch.cell(row=r_idx, column=col_idx)
                c.font = font_body
                c.border = cell_border
                c.fill = c_fill_defaut
                
                if col_idx in [1, 4, 7, 9]:
                    c.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx == 6:
                    c.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    c.alignment = Alignment(horizontal="left", vertical="center")
                
                if col_idx == 4:
                    if "Présent" in statut_aff:
                        c.fill = PatternFill(start_color=VERT_BG, end_color=VERT_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color=VERT_TXT)
                    elif "1/2" in statut_aff:
                        c.fill = PatternFill(start_color=JAUNE_BG, end_color=JAUNE_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color=JAUNE_TXT)
                    elif "Autorisée" in statut_aff:
                        c.fill = PatternFill(start_color=ORANGE_BG, end_color=ORANGE_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color=ORANGE_TXT)
                    elif "Absence" in statut_aff:
                        c.fill = PatternFill(start_color=ROUGE_BG, end_color=ROUGE_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color=ROUGE_TXT)
                        
                elif col_idx == 7:
                    if "Conforme" in qual_clean and "Non" not in qual_clean:
                        c.fill = PatternFill(start_color=VERT_BG, end_color=VERT_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color=VERT_TXT)
                    elif "Acceptable" in qual_clean:
                        c.fill = PatternFill(start_color=JAUNE_BG, end_color=JAUNE_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color=JAUNE_TXT)
                    elif "Non conforme" in qual_clean:
                        c.fill = PatternFill(start_color=ROUGE_BG, end_color=ROUGE_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, bold=True, color=ROUGE_TXT)
                        
                elif col_idx == 8:
                    if "bon rendement" in obs_clean or "Excellent" in obs_clean or "Bon travail" in obs_clean:
                        c.fill = PatternFill(start_color=VERT_BG, end_color=VERT_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, color=VERT_TXT)
                    elif "Rendement moyen" in obs_clean or "Moyen" in obs_clean:
                        c.fill = PatternFill(start_color=JAUNE_BG, end_color=JAUNE_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, color=JAUNE_TXT)
                    elif "Faible rendement" in obs_clean or "surveiller" in obs_clean:
                        c.fill = PatternFill(start_color=ROUGE_BG, end_color=ROUGE_BG, fill_type="solid")
                        c.font = Font(name="Segoe UI", size=9, color=ROUGE_TXT)
                        
            r_idx += 1
            
        for col_lettre, larg in largeurs_detail.items():
            ws_ch.column_dimensions[col_lettre].width = larg

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if "sync_notif" not in st.session_state:
    st.session_state["sync_notif"] = None

if st.session_state["sync_notif"]:
    st.success(st.session_state["sync_notif"])

st.title("🏗️️ Suivi Chantier & Étanchéité")

dict_conducteurs = get_conducteurs_dict()
nom_c1 = dict_conducteurs.get("c1", "Conducteur 1")
nom_c2 = dict_conducteurs.get("c2", "Conducteur 2")

menu_general = st.radio(
    "Espace de travail",
    [f"👷 Espace {nom_c1}", f"👷 Espace {nom_c2}", "🔐 Espace Admin (Direction)"],
    horizontal=False
)

def interface_saisie_conducteur(conducteur_id_tag, default_nom):
    photo_cond = get_photo_path(conducteur_id_tag)
    
    col_h1, col_h2 = st.columns([1, 4])
    with col_h1:
        if photo_cond:
            st.image(photo_cond, width=65)
        else:
            st.markdown("<div style='font-size:40px;text-align:center;'>👷</div>", unsafe_allow_html=True)
    with col_h2:
        st.subheader(f"Pointage — {default_nom}")

    date_choisie = st.date_input("📅 Date de saisie", value=date.today(), key=f"date_{conducteur_id_tag}")
    date_str = str(date_choisie)
    chantiers_autorises = get_chantiers_conducteur(conducteur_id_tag)

    if not chantiers_autorises:
        st.warning(f"⚠️ Aucun chantier attribué à {default_nom}.\n\nVeuillez contacter l'administrateur.")
        return

    chantiers_non_faits = [ch for ch in chantiers_autorises if not est_deja_valide(date_str, ch)]
    chantiers_faits = [ch for ch in chantiers_autorises if est_deja_valide(date_str, ch)]

    if not chantiers_non_faits:
        st.success(f"🎉 Tous vos chantiers du {date_str} ont été pointés et validés avec succès !")
        st.markdown(f"**Chantiers complétés :** `{', '.join(chantiers_faits)}`")
        
        st.markdown("---")
        afficher_modif = st.checkbox("Modifier un chantier déjà validé", key=f"cb_rev_{conducteur_id_tag}")
        if not afficher_modif:
            return
        liste_options_chantiers = chantiers_faits
    else:
        liste_options_chantiers = chantiers_non_faits

    chantier_choisi = st.selectbox("📍 Sélectionner le Chantier à pointer", liste_options_chantiers, key=f"ch_sel_{conducteur_id_tag}")
    deja_fait = est_deja_valide(date_str, chantier_choisi)

    if deja_fait:
        st.info(f"🟢 **Pointage déjà enregistré pour {chantier_choisi} le {date_str}.**")

    df_w = get_workers_df()
    equipe_active = df_w[df_w["chantier_fixe"] == chantier_choisi]

    st.markdown(f"**Effectif sur {chantier_choisi} :** `{len(equipe_active)}` ouvrier(s)")

    if equipe_active.empty:
        st.warning(f"⚠️ Aucun ouvrier affecté à {chantier_choisi}.")
        return

    step_key = f"step_{conducteur_id_tag}_{chantier_choisi}"
    if step_key not in st.session_state:
        st.session_state[step_key] = 1

    presence_data_key = f"presence_cache_{conducteur_id_tag}_{chantier_choisi}"
    if presence_data_key not in st.session_state:
        st.session_state[presence_data_key] = {}

    # ÉTAPE 1 : POINTAGE DANS LE CADRE UNIQUE PAR OUVRIER
    if st.session_state[step_key] == 1:
        st.markdown("### 📋 Étape 1 : Présence & Absences")
        st.caption("Sélectionnez le statut de chaque ouvrier puis passez à l'étape suivante.")

        presence_temp = {}
        for _, row in equipe_active.iterrows():
            w_id = row['id']
            w_nom = row['nom']

            # Conteneur unifié
            with st.container():
                st.markdown(f"""
                <div class='worker-frame'>
                    <div class='worker-header'>
                        <div>{get_avatar_html(w_nom)}</div>
                        <div>
                            <h3 style='margin:0; font-size:1.15rem; color:#F9FAFB;'>{w_nom}</h3>
                            <span style='font-size:0.8rem; color:#9CA3AF;'>ID: #{w_id} • {chantier_choisi}</span>
                        </div>
                    </div>
                """, unsafe_allow_html=True)

                statut_val = st.selectbox(
                    "Statut de présence :",
                    [
                        "Présent",
                        "1/2 journée",
                        "Absence Autorisée",
                        "Absence Non Autorisée (Injustifiée)"
                    ],
                    key=f"st1_{conducteur_id_tag}_{w_id}"
                )

                motif_abs = ""
                if "Absence" in statut_val:
                    motif_abs = st.text_input("Motif de l'absence :", placeholder="Ex: congé, maladie, arrêt...", key=f"abs_m_{conducteur_id_tag}_{w_id}")

                st.markdown("</div>", unsafe_allow_html=True)

            presence_temp[w_id] = {
                "nom": w_nom,
                "statut": statut_val,
                "motif_absence": motif_abs
            }

        if st.button("➡️ Étape 2 : Production & Travaux des Présents", type="primary", use_container_width=True):
            st.session_state[presence_data_key] = presence_temp
            st.session_state[step_key] = 2
            st.rerun()

    # ÉTAPE 2 : PRODUCTION & CONTRÔLE DANS LE CADRE UNIQUE PAR OUVRIER
    elif st.session_state[step_key] == 2:
        st.markdown("### 🔨 Étape 2 : Production & Appréciation")
        st.caption("Saisie uniquement pour les présents. Les absences sont validées directement.")

        if st.button("⬅️ Retour au Pointage (Étape 1)", use_container_width=True):
            st.session_state[step_key] = 1
            st.rerun()

        presence_enregistree = st.session_state.get(presence_data_key, {})
        toutes_les_taches = get_all_taches()
        donnees_finales = {}

        for _, row in equipe_active.iterrows():
            w_id = row['id']
            w_nom = row['nom']
            infos_p = presence_enregistree.get(w_id, {"statut": "Présent", "motif_absence": ""})
            st_val = infos_p["statut"]

            if "Présent" in st_val or "1/2" in st_val:
                with st.container():
                    st.markdown(f"""
                    <div class='worker-frame'>
                        <div class='worker-header'>
                            <div>{get_avatar_html(w_nom)}</div>
                            <div>
                                <h3 style='margin:0; font-size:1.15rem; color:#F9FAFB;'>{w_nom}</h3>
                                <span style='font-size:0.85rem; color:#10B981; font-weight:600;'>🟢 Statut : {st_val}</span>
                            </div>
                        </div>
                    """, unsafe_allow_html=True)

                    tache_val = st.selectbox("Corps d'état / Tâche :", toutes_les_taches, key=f"tch2_{conducteur_id_tag}_{w_id}")
                    est_bricol_defaut = ("BRICOL" in tache_val.upper()) or (tache_val in ["DIVERS", "nettoyage", "PONSAGE"])
                    type_travail = st.selectbox("Type d'activité :", ["Métrage (m² / ml)", "Bricolage / Sans métrage"], index=1 if est_bricol_defaut else 0, key=f"typ2_{conducteur_id_tag}_{w_id}")

                    if type_travail == "Métrage (m² / ml)":
                        qte_val = st.number_input("Métré réalisé :", min_value=0.0, step=1.0, value=25.0, key=f"qte2_{conducteur_id_tag}_{w_id}")
                        unite_val = "m²"
                    else:
                        unite_val = "Sans métrage"
                        qte_val = 1.0

                    c_ev1, c_ev2 = st.columns(2)
                    with c_ev1:
                        eval_ouvrier = st.selectbox(
                            "Rendement / Implication :",
                            [
                                "Très bon rendement (Excellent)",
                                "Bon travail (Régulier)",
                                "Rendement moyen (Moyen)",
                                "Faible rendement (À surveiller)"
                            ],
                            key=f"eval_kh_{conducteur_id_tag}_{w_id}"
                        )
                    with c_ev2:
                        apprec_qualite = st.selectbox(
                            "Qualité d'exécution :",
                            [
                                "Conforme / Soigné",
                                "Acceptable",
                                "Non conforme"
                            ],
                            key=f"qual_{conducteur_id_tag}_{w_id}"
                        )

                    obs_val = st.text_input("Observation libre :", placeholder="Ex: terrasse sud, acrotères...", key=f"obs2_{conducteur_id_tag}_{w_id}")
                    st.markdown("</div>", unsafe_allow_html=True)

                donnees_finales[w_id] = {
                    "statut": st_val,
                    "tache": tache_val,
                    "quantite": qte_val,
                    "unite": unite_val,
                    "appreciation": apprec_qualite,
                    "observation": f"{eval_ouvrier} | {obs_val}" if obs_val else eval_ouvrier,
                    "eval_travailleur": eval_ouvrier
                }
            else:
                donnees_finales[w_id] = {
                    "statut": st_val,
                    "tache": "-",
                    "quantite": 0.0,
                    "unite": "-",
                    "appreciation": "-",
                    "observation": infos_p.get("motif_absence", ""),
                    "eval_travailleur": "-"
                }

        label_save = "🔄 Mettre à jour la validation" if deja_fait else "💾 Valider définitivement le chantier"
        if st.button(label_save, type="primary", use_container_width=True):
            conn = get_db_connection()
            c = conn.cursor()

            for w_id, d in donnees_finales.items():
                st_val = d["statut"]
                qte = d["quantite"]
                unite = d["unite"]
                app = d["appreciation"]
                eval_w = d.get("eval_travailleur", "-")

                if "Présent" in st_val:
                    base_p = 35.0
                    pts_prod = 35.0 if (unite == "Sans métrage" or qte >= 30) else (25.0 if qte >= 20 else 15.0)
                    pts_app = 20.0 if "Conforme" in app and "Non" not in app else (10.0 if "Acceptable" in app else 0.0)
                    pts_w = 10.0 if "Excellent" in eval_w else (8.0 if "Régulier" in eval_w else (4.0 if "Moyen" in eval_w else 0.0))
                    score = min(base_p + pts_prod + pts_app + pts_w, 100.0)
                elif "1/2" in st_val:
                    score = 35.0
                elif "Autorisée" in st_val:
                    score = None
                else:
                    score = 0.0

                c.execute("""
                    INSERT OR REPLACE INTO pointages (date_jour, chantier, conducteur, worker_id, statut, tache, quantite, unite, appreciation, observation, score)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, (date_str, chantier_choisi, default_nom, w_id, st_val, d['tache'], qte, unite, app, d['observation'], score))

            conn.commit()
            conn.close()

            st.session_state[step_key] = 1
            st.session_state["sync_notif"] = f"✅ Pointage validé avec succès pour {chantier_choisi} !"
            st.rerun()

if menu_general == f"👷 Espace {nom_c1}":
    interface_saisie_conducteur("c1", nom_c1)

elif menu_general == f"👷 Espace {nom_c2}":
    interface_saisie_conducteur("c2", nom_c2)

elif menu_general == "🔐 Espace Admin (Direction)":
    st.subheader("Accès Sécurisé - Administration")

    if "admin_logged_in" not in st.session_state:
        st.session_state["admin_logged_in"] = False

    if not st.session_state["admin_logged_in"]:
        mdp = st.text_input("Mot de passe Administrateur", type="password", placeholder="Entrez le mot de passe...")
        if st.button("Connexion", type="primary", use_container_width=True):
            if mdp == ADMIN_PASSWORD:
                st.session_state["admin_logged_in"] = True
                st.rerun()
            else:
                st.error("Mot de passe incorrect.")
    else:
        if st.button("🚪 Déconnexion", use_container_width=True):
            st.session_state["admin_logged_in"] = False
            st.session_state["admin_active_module"] = None
            st.rerun()

        conn = get_db_connection()
        total_p = conn.cursor().execute("SELECT COUNT(*) FROM pointages").fetchone()[0]
        conn.close()
        st.info(f"📊 Registre : **{total_p} pointages enregistrés**")

        if "admin_active_module" not in st.session_state:
            st.session_state["admin_active_module"] = None

        MODULES_ADMIN = [
            ("mod_rapport", "📊 Bilan Mensuel & Rapports"),
            ("mod_corriger", "✏️ Corriger un Pointage"),
            ("mod_transfert", "🔄 Transférer un Ouvrier"),
            ("mod_equipes", "⚡ Équipes par Chantier"),
            ("mod_chantiers_cond", "👷 Affecter les Chantiers"),
            ("mod_profils", "👥 Profils & Photos"),
            ("mod_chantiers_taches", "🏗️ Chantiers & Tâches"),
        ]

        if st.session_state["admin_active_module"] is None:
            st.markdown("### 🎛 Menu Administrateur")
            st.markdown("##### ⚡ Actions Quotidiennes")
            for tag, titre in MODULES_ADMIN[:3]:
                if st.button(titre, key=f"btn_case_{tag}", use_container_width=True):
                    st.session_state["admin_active_module"] = tag
                    st.rerun()

            st.markdown("##### ⚙️ Configurations & Éditions")
            for tag, titre in MODULES_ADMIN[3:]:
                if st.button(titre, key=f"btn_case_{tag}", use_container_width=True):
                    st.session_state["admin_active_module"] = tag
                    st.rerun()

        else:
            mod_actuel = st.session_state["admin_active_module"]

            if st.button("⬅️ Retour au Menu Admin", type="secondary", use_container_width=True):
                st.session_state["admin_active_module"] = None
                st.rerun()

            st.markdown("---")

            if mod_actuel == "mod_rapport":
                st.markdown("### 📊 Bilan Mensuel & Rapports")

                conn = get_db_connection()
                query_admin = """
                    SELECT 
                        p.id AS ID,
                        p.date_jour AS [Date],
                        p.chantier AS [Chantier],
                        COALESCE(p.conducteur, 'Non spécifié') AS [Conducteur],
                        COALESCE(w.nom, 'Ouvrier #' || p.worker_id) AS [Ouvrier],
                        p.statut AS [Statut],
                        p.tache AS [Tâche],
                        p.quantite AS [Quantite],
                        p.unite AS [Unite],
                        COALESCE(p.appreciation, '-') AS [Qualité],
                        COALESCE(p.observation, '-') AS [Observation],
                        p.score AS [Score]
                    FROM pointages p
                    LEFT JOIN workers w ON p.worker_id = w.id
                    ORDER BY p.chantier ASC, p.date_jour DESC, p.id DESC
                """
                df_all = pd.read_sql_query(query_admin, conn)
                conn.close()

                if df_all.empty:
                    st.warning("⚠️ Aucune donnée enregistrée.")
                else:
                    df_all["Mois_Annee"] = df_all["Date"].str.slice(0, 7)
                    mois_dispos = sorted(df_all["Mois_Annee"].unique().tolist(), reverse=True)

                    c_m1, c_m2 = st.columns([1, 2])
                    with c_m1:
                        mois_choisi = st.selectbox("📅 Mois :", mois_dispos)
                    with c_m2:
                        filtre_ch = st.selectbox("📍 Chantier :", ["Tous les chantiers"] + sorted(df_all["Chantier"].unique().tolist()))

                    df_mois_actuel = df_all[df_all["Mois_Annee"] == mois_choisi].copy()
                    if filtre_ch != "Tous les chantiers":
                        df_mois_actuel = df_mois_actuel[df_mois_actuel["Chantier"] == filtre_ch]

                    total_ouv_actifs = df_mois_actuel["Ouvrier"].nunique()
                    p_c = sum(1 for v in df_mois_actuel["Statut"] if "Présent" in str(v))
                    d_c = sum(1 for v in df_mois_actuel["Statut"] if "1/2" in str(v))
                    j_payes = p_c + (d_c * 0.5)
                    m2_tot = df_mois_actuel[(df_mois_actuel["Unite"] == "m²") & (df_mois_actuel["Quantite"] > 0)]["Quantite"].sum()

                    k1, k2 = st.columns(2)
                    k1.metric("Ouvriers Actifs", f"{total_ouv_actifs}")
                    k2.metric("Jours Payés", f"{j_payes} j")

                    k3, k4 = st.columns(2)
                    k3.metric("Absences", f"{sum(1 for v in df_mois_actuel['Statut'] if 'Absence' in str(v))}")
                    k4.metric("Production (m²)", f"{m2_tot:.1f} m²")

                    st.markdown("---")
                    tab_synth1, tab_synth2, tab_synth3 = st.tabs([
                        "🏗️ Synthèse par Chantier",
                        "👥 Bilan & Primes Ouvriers",
                        "📋 Registre Détaillé"
                    ])

                    with tab_synth1:
                        ch_synth = []
                        for ch_name, grp in df_mois_actuel.groupby("Chantier"):
                            p_cnt = sum(1 for v in grp["Statut"] if "Présent" in str(v))
                            d_cnt = sum(1 for v in grp["Statut"] if "1/2" in str(v))
                            a_cnt = sum(1 for v in grp["Statut"] if "Absence" in str(v))
                            j_val = p_cnt + (d_cnt * 0.5)
                            m2_val = grp[(grp["Unite"] == "m²") & (grp["Quantite"] > 0)]["Quantite"].sum()

                            ch_synth.append({
                                "Chantier": ch_name,
                                "Effectif": grp["Ouvrier"].nunique(),
                                "Présents (J)": p_cnt,
                                "1/2 J": d_cnt,
                                "Absences": a_cnt,
                                "Jours Payés": f"{j_val} j",
                                "Métré (m²)": f"{m2_val:.1f} m²"
                            })

                        df_ch_synth = pd.DataFrame(ch_synth)
                        st.dataframe(df_ch_synth, use_container_width=True, hide_index=True)

                    with tab_synth2:
                        st.markdown("##### 👥 Performance Individuelle & Éligibilité Primes")
                        ouv_synth_list = []
                        for w_name, w_grp in df_mois_actuel.groupby("Ouvrier"):
                            p_c = sum(1 for v in w_grp["Statut"] if "Présent" in str(v))
                            d_c = sum(1 for v in w_grp["Statut"] if "1/2" in str(v))
                            j_p = p_c + (d_c * 0.5)
                            abs_injust = sum(1 for v in w_grp["Statut"] if "Non Autorisée" in str(v) or "Injustifiée" in str(v))
                            prod_w = w_grp[(w_grp["Unite"] == "m²") & (w_grp["Quantite"] > 0)]["Quantite"].sum()
                            
                            scores_v = pd.to_numeric(w_grp["Score"], errors='coerce').dropna()
                            sc_m = scores_v.mean() if not scores_v.empty else 0.0

                            if abs_injust > 0:
                                mention = "❌ Non Éligible (Absence)"
                            elif sc_m >= 90.0 and j_p >= 20.0:
                                mention = "🥇 Prime Maximale (A+)"
                            elif sc_m >= 80.0:
                                mention = "🥈 Prime Standard (A)"
                            elif sc_m >= 70.0:
                                mention = "🥉 Encouragement (B)"
                            else:
                                mention = "Sans Prime"

                            ouv_synth_list.append({
                                "Ouvrier": w_name,
                                "Chantier": w_grp["Chantier"].mode()[0] if not w_grp.empty else "-",
                                "Jours Payés": f"{j_p} j",
                                "Absences Injustifiées": abs_injust,
                                "Production (m²)": f"{prod_w:.1f} m²",
                                "Score (/100)": f"{sc_m:.1f}",
                                "Statut Prime": mention
                            })

                        df_ouv_synth = pd.DataFrame(ouv_synth_list)
                        st.dataframe(df_ouv_synth, use_container_width=True, hide_index=True)

                    with tab_synth3:
                        df_vue_detail = df_mois_actuel.copy()
                        df_vue_detail["Production"] = df_vue_detail.apply(
                            lambda r: "Bricolage" if r["Unite"] == "Sans métrage" else (f"{r['Quantite']} {r['Unite']}" if r["Quantite"] > 0 else "-"),
                            axis=1
                        )
                        cols_finales = ["Chantier", "Date", "Conducteur", "Ouvrier", "Statut", "Tâche", "Production", "Qualité", "Score"]
                        st.dataframe(df_vue_detail[cols_finales], use_container_width=True, hide_index=True)

                    st.markdown("#### 📥 Téléchargements")
                    excel_pro_bytes = generer_classeur_pro_excel(df_mois_actuel, mois_choisi)
                    st.download_button(
                        label="📗 Télécharger le Rapport Excel (.xlsx) avec Feuille Primes & Chantiers",
                        data=excel_pro_bytes,
                        file_name=f"rapport_chantiers_{mois_choisi}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                        use_container_width=True
                    )

            elif mod_actuel == "mod_corriger":
                st.markdown("### ✏️ Corriger un Pointage")
                
                conn = get_db_connection()
                query_admin = """
                    SELECT 
                        p.id AS ID,
                        p.date_jour AS [Date],
                        p.chantier AS [Chantier],
                        COALESCE(p.conducteur, 'Non spécifié') AS [Conducteur],
                        COALESCE(w.nom, 'Ouvrier #' || p.worker_id) AS [Ouvrier],
                        p.statut AS [Statut],
                        p.tache AS [Tâche],
                        p.quantite AS [Quantite],
                        p.unite AS [Unite],
                        COALESCE(p.appreciation, '-') AS [Qualité],
                        COALESCE(p.observation, '-') AS [Observation]
                    FROM pointages p
                    LEFT JOIN workers w ON p.worker_id = w.id
                    ORDER BY p.date_jour DESC, p.id DESC
                """
                df_pts = pd.read_sql_query(query_admin, conn)
                conn.close()

                if df_pts.empty:
                    st.info("Aucun pointage à corriger pour le moment.")
                else:
                    df_pts["label"] = df_pts.apply(
                        lambda r: f"ID #{r['ID']} | {r['Date']} | {r['Chantier']} | {r['Ouvrier']}",
                        axis=1
                    )
                    ligne_choisie = st.selectbox("Sélectionner la ligne :", df_pts["label"].tolist(), key="sel_pt_correction_prio")
                    row_sel = df_pts[df_pts["label"] == ligne_choisie].iloc[0]
                    pt_id = int(row_sel["ID"])

                    st.info(f"**{row_sel['Ouvrier']}** — {row_sel['Chantier']} ({row_sel['Date']})")

                    statuts_possibles = ["Présent", "1/2 journée", "Absence Autorisée", "Absence Non Autorisée (Injustifiée)"]
                    idx_st = statuts_possibles.index(row_sel["Statut"]) if row_sel["Statut"] in statuts_possibles else 0
                    mod_statut = st.selectbox("Statut :", statuts_possibles, index=idx_st, key=f"mod_st_{pt_id}")
                    
                    toutes_les_taches_dispos = get_all_taches()
                    idx_tch = toutes_les_taches_dispos.index(row_sel["Tâche"]) if row_sel["Tâche"] in toutes_les_taches_dispos else 0
                    mod_tache = st.selectbox("Tâche :", toutes_les_taches_dispos, index=idx_tch, key=f"mod_tch_{pt_id}")

                    type_activite_mod = st.selectbox("Type :", ["Métrage (m² / ml)", "Bricolage / Sans métrage"], index=1 if row_sel["Unite"] == "Sans métrage" else 0, key=f"mod_typ_{pt_id}")
                    if type_activite_mod == "Métrage (m² / ml)":
                        mod_qte = st.number_input("Quantité :", min_value=0.0, step=1.0, value=float(row_sel["Quantite"]), key=f"mod_qte_{pt_id}")
                        mod_unite = "m²"
                    else:
                        mod_qte = 1.0
                        mod_unite = "Sans métrage"

                    qualites_possibles = ["Conforme / Soigné", "Acceptable", "Non conforme"]
                    clean_qual_existante = str(row_sel["Qualité"]).replace("/ À reprendre", "").replace("🟢", "").replace("🟡", "").replace("🔴", "").strip()
                    idx_qual = 0
                    for idx_q, q_label in enumerate(qualites_possibles):
                        if clean_qual_existante in q_label or q_label in clean_qual_existante:
                            idx_qual = idx_q
                            break
                    mod_apprec = st.selectbox("Qualité :", qualites_possibles, index=idx_qual, key=f"mod_qual_{pt_id}")

                    mod_obs = st.text_input("Observation :", value=str(row_sel["Observation"]) if row_sel["Observation"] != "-" else "", key=f"mod_obs_{pt_id}")

                    col_b1, col_b2 = st.columns(2)
                    with col_b1:
                        if st.button("💾 Enregistrer Correction", type="primary", use_container_width=True):
                            if "Présent" in mod_statut:
                                base_p = 40.0
                                pts_prod = 40.0 if (mod_unite == "Sans métrage" or mod_qte >= 30) else (30.0 if mod_qte >= 20 else 15.0)
                                pts_app = 20.0 if "Conforme" in mod_apprec and "Non" not in mod_apprec else (10.0 if "Acceptable" in mod_apprec else 0.0)
                                score_corr = min(base_p + pts_prod + pts_app, 100.0)
                            elif "1/2" in mod_statut:
                                score_corr = 35.0
                            elif "Autorisée" in mod_statut:
                                score_corr = None
                            else:
                                score_corr = 0.0

                            conn = get_db_connection()
                            c = conn.cursor()
                            c.execute("""
                                UPDATE pointages 
                                SET statut = ?, tache = ?, quantite = ?, unite = ?, appreciation = ?, observation = ?, score = ?
                                WHERE id = ?
                            """, (mod_statut, mod_tache, mod_qte, mod_unite, mod_apprec, mod_obs, score_corr, pt_id))
                            conn.commit()
                            conn.close()
                            st.session_state["sync_notif"] = f"✅ Pointage #{pt_id} rectifié !"
                            st.rerun()

                    with col_b2:
                        if st.button("🗑️ Supprimer Pointage", type="secondary", use_container_width=True):
                            conn = get_db_connection()
                            c = conn.cursor()
                            c.execute("DELETE FROM pointages WHERE id = ?", (pt_id,))
                            conn.commit()
                            conn.close()
                            st.session_state["sync_notif"] = f"🗑 Pointage #{pt_id} supprimé."
                            st.rerun()

            elif mod_actuel == "mod_transfert":
                st.markdown("### 🔄 Transférer un Ouvrier")
                df_w_tr = get_workers_df()
                ouv_sel = st.selectbox("Ouvrier à transférer :", df_w_tr["nom"].tolist(), key="tr_o_prio")
                infos = df_w_tr[df_w_tr["nom"] == ouv_sel].iloc[0]
                st.info(f"Chantier actuel : **{infos['chantier_fixe']}**")
                dest_ch = st.selectbox("Vers chantier :", [c for c in get_all_chantiers() if c != infos["chantier_fixe"]], key="dest_ch_prio")

                if st.button("Confirmer le Transfert", type="primary", use_container_width=True):
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute("UPDATE workers SET chantier_fixe = ? WHERE id = ?", (dest_ch, int(infos["id"])))
                    conn.commit()
                    conn.close()
                    st.session_state["sync_notif"] = f"🔄 {ouv_sel} transféré vers {dest_ch} !"
                    st.rerun()

            elif mod_actuel == "mod_equipes":
                st.markdown("### ⚡ Équipes par Chantier")
                ch_dispos_actifs = [c for c in get_all_chantiers() if c != "EN ATTENTE / DEPOT"]
                ch_cible = st.selectbox("Chantier à configurer :", ch_dispos_actifs, key="ch_eq_cfg")
                
                df_w_eq = get_workers_df()
                actuels = df_w_eq[df_w_eq["chantier_fixe"] == ch_cible]["nom"].tolist()
                ouvriers_libres = df_w_eq[df_w_eq["chantier_fixe"] == "EN ATTENTE / DEPOT"]["nom"].tolist()
                options_autorisees = sorted(list(set(actuels + ouvriers_libres)))
                
                st.caption(f"💡 `{len(ouvriers_libres)}` ouvrier(s) libre(s). Les ouvriers déjà affectés ailleurs sont masqués.")
                
                nouv_eq = st.multiselect(
                    f"Ouvriers sur {ch_cible} :",
                    options=options_autorisees,
                    default=actuels,
                    key=f"ms_eq_{ch_cible}_cfg"
                )

                if st.button("💾 Verrouiller l'Équipe", type="primary", use_container_width=True):
                    conn = get_db_connection()
                    c = conn.cursor()
                    for nom in actuels:
                        if nom not in nouv_eq:
                            c.execute("UPDATE workers SET chantier_fixe = 'EN ATTENTE / DEPOT' WHERE nom = ?", (nom,))
                    for nom in nouv_eq:
                        c.execute("UPDATE workers SET chantier_fixe = ? WHERE nom = ?", (ch_cible, nom))
                    conn.commit()
                    conn.close()
                    st.session_state["sync_notif"] = f"🔄 Équipe de {ch_cible} enregistrée ({len(nouv_eq)} ouvriers) !"
                    st.rerun()

            elif mod_actuel == "mod_chantiers_cond":
                st.markdown("### 👷 Affecter les Chantiers")
                ch_dispos = [c for c in get_all_chantiers() if c != "EN ATTENTE / DEPOT"]
                
                st.markdown(f"**{nom_c1}**")
                act_c1 = get_chantiers_conducteur("c1")
                nouv_c1 = st.multiselect("Chantiers c1 :", options=ch_dispos, default=act_c1, key="ms_c1_cfg")

                st.markdown(f"**{nom_c2}**")
                act_c2 = get_chantiers_conducteur("c2")
                nouv_c2 = st.multiselect("Chantiers c2 :", options=ch_dispos, default=act_c2, key="ms_c2_cfg")

                if st.button("💾 Sauvegarder les Affectations", type="primary", use_container_width=True):
                    conn = get_db_connection()
                    c = conn.cursor()
                    c.execute("DELETE FROM conducteur_chantiers WHERE conducteur_tag = 'c1'")
                    for ch in nouv_c1:
                        c.execute("INSERT INTO conducteur_chantiers (conducteur_tag, chantier) VALUES ('c1', ?)", (ch,))
                    c.execute("DELETE FROM conducteur_chantiers WHERE conducteur_tag = 'c2'")
                    for ch in nouv_c2:
                        c.execute("INSERT INTO conducteur_chantiers (conducteur_tag, chantier) VALUES ('c2', ?)", (ch,))
                    conn.commit()
                    conn.close()
                    st.session_state["sync_notif"] = "🔄 Affectations sauvegardées !"
                    st.rerun()

            elif mod_actuel == "mod_profils":
                st.markdown("### 👥 Profils & Photos")
                tab_cond, tab_ouv = st.tabs(["👷 Conducteurs", "👷 Ouvriers"])

                with tab_cond:
                    c_tag = st.radio("Conducteur :", ["c1", "c2"], format_func=lambda x: f"Conducteur 1 ({nom_c1})" if x == 'c1' else f"Conducteur 2 ({nom_c2})")
                    n_act = nom_c1 if c_tag == 'c1' else nom_c2
                    ph_c = get_photo_path(c_tag)
                    if ph_c:
                        st.image(ph_c, width=120)
                    n_nom_c = st.text_input("Nom affiché :", value=n_act, key=f"in_nc_{c_tag}_cfg")
                    n_ph_c = st.file_uploader("Photo profil :", type=["jpg", "jpeg", "png"], key=f"up_pc_{c_tag}_cfg")
                    if st.button("💾 Enregistrer Profil Conducteur", type="primary", use_container_width=True):
                        n_propre = n_nom_c.strip() or ("Conducteur 1" if c_tag == 'c1' else "Conducteur 2")
                        conn = get_db_connection()
                        c = conn.cursor()
                        c.execute("UPDATE conducteurs_meta SET nom_affiche = ? WHERE tag = ?", (n_propre, c_tag))
                        conn.commit()
                        conn.close()

                        if n_ph_c is not None:
                            ext = n_ph_c.name.split(".")[-1].lower()
                            nom_f = f"{c_tag}.{ext}"
                            for o_ext in [".jpg", ".jpeg", ".png"]:
                                old_path = os.path.join(PHOTOS_DIR, f"{c_tag}{o_ext}")
                                if os.path.exists(old_path):
                                    os.remove(old_path)
                            Image.open(n_ph_c).save(os.path.join(PHOTOS_DIR, nom_f))

                        st.session_state["sync_notif"] = f"✅ Conducteur {n_propre} mis à jour !"
                        st.rerun()

                with tab_ouv:
                    df_w_m = get_workers_df()
                    sub_ed, sub_ad, sub_dl = st.tabs(["✏ Modifier", "➕ Ajouter", "🗑 Supprimer"])
                    with sub_ed:
                        o_sel = st.selectbox("Sélectionner l'ouvrier :", df_w_m["nom"].tolist(), key="sel_ouv_cfg")
                        ph_o = get_photo_path(o_sel)
                        if ph_o:
                            st.image(ph_o, width=120)
                        n_nom_o = st.text_input("Nom complet :", value=o_sel, key=f"in_no_{o_sel}_cfg")
                        n_ph_o = st.file_uploader("Photo profil ouvrier :", type=["jpg", "jpeg", "png"], key=f"up_po_{o_sel}_cfg")

                        if st.button("💾 Enregistrer Modifications", type="primary", use_container_width=True):
                            nom_p = n_nom_o.strip()
                            conn = get_db_connection()
                            c = conn.cursor()
                            if nom_p and nom_p != o_sel:
                                try:
                                    c.execute("UPDATE workers SET nom = ? WHERE nom = ?", (nom_p, o_sel))
                                    conn.commit()
                                    o_cl = o_sel.replace(" ", "_")
                                    n_cl = nom_p.replace(" ", "_")
                                    for ext in [".jpg", ".jpeg", ".png"]:
                                        if os.path.exists(os.path.join(PHOTOS_DIR, f"{o_cl}{ext}")):
                                            os.rename(os.path.join(PHOTOS_DIR, f"{o_cl}{ext}"), os.path.join(PHOTOS_DIR, f"{n_cl}{ext}"))
                                    nom_ref = nom_p
                                except sqlite3.IntegrityError:
                                    conn.close()
                                    st.error("Ce nom existe déjà.")
                                    st.stop()
                            else:
                                nom_ref = o_sel

                            if n_ph_o is not None:
                                ext = n_ph_o.name.split(".")[-1].lower()
                                nom_f = f"{nom_ref.replace(' ', '_')}.{ext}"
                                Image.open(n_ph_o).save(os.path.join(PHOTOS_DIR, nom_f))

                            conn.close()
                            st.session_state["sync_notif"] = f"✅ Profil {nom_ref} mis à jour !"
                            st.rerun()

                    with sub_ad:
                        nom_nouv = st.text_input("Nom et prénom :", key="in_add_w_nom_cfg")
                        ch_init = st.selectbox("Chantier initial :", get_all_chantiers(), key="in_add_w_ch_cfg")
                        ph_nouv = st.file_uploader("Photo profil :", type=["jpg", "jpeg", "png"], key="in_add_w_ph_cfg")
                        if st.button("➕ Ajouter l'Ouvrier", type="primary", use_container_width=True):
                            n_net = nom_nouv.strip()
                            if n_net:
                                conn = get_db_connection()
                                c = conn.cursor()
                                try:
                                    c.execute("INSERT INTO workers (nom, chantier_fixe) VALUES (?, ?)", (n_net, ch_init))
                                    conn.commit()
                                    conn.close()
                                    if ph_nouv is not None:
                                        ext = ph_nouv.name.split(".")[-1].lower()
                                        Image.open(ph_nouv).save(os.path.join(PHOTOS_DIR, f"{n_net.replace(' ', '_')}.{ext}"))
                                    st.session_state["sync_notif"] = f"✅ {n_net} ajouté !"
                                    st.rerun()
                                except sqlite3.IntegrityError:
                                    conn.close()
                                    st.error("Cet ouvrier existe déjà.")

                    with sub_dl:
                        ouv_del_s = st.selectbox("Ouvrier à supprimer :", df_w_m["nom"].tolist(), key="sel_del_ouv_cfg")
                        if st.button("❌ Supprimer Définitivement", type="secondary", use_container_width=True):
                            conn = get_db_connection()
                            c = conn.cursor()
                            c.execute("DELETE FROM workers WHERE nom = ?", (ouv_del_s,))
                            conn.commit()
                            conn.close()
                            st.session_state["sync_notif"] = f"🗑️ {ouv_del_s} supprimé de la base."
                            st.rerun()

            elif mod_actuel == "mod_chantiers_taches":
                st.markdown("### 🏗️ Chantiers & Tâches")
                tab_ch, tab_tch = st.tabs(["📍 Chantiers", "🔨 Tâches"])

                with tab_ch:
                    liste_ch = get_all_chantiers()
                    nouveau_ch = st.text_input("Nouveau chantier :", key="in_add_ch_cfg")
                    if st.button("➕ Ajouter Chantier", type="primary", use_container_width=True):
                        n_c = nouveau_ch.strip()
                        if n_c:
                            conn = get_db_connection()
                            c = conn.cursor()
                            try:
                                c.execute("INSERT INTO chantiers_ref (nom) VALUES (?)", (n_c,))
                                conn.commit()
                                conn.close()
                                st.session_state["sync_notif"] = f"✅ Chantier {n_c} ajouté !"
                                st.rerun()
                            except sqlite3.IntegrityError:
                                conn.close()
                                st.error("Ce chantier existe déjà.")

                    st.markdown("---")
                    ch_a_renom = st.selectbox("Chantier à renommer :", liste_ch, key="sel_mod_ch_cfg")
                    ch_new_n = st.text_input("Nouveau libellé :", value=ch_a_renom, key="in_renom_ch_cfg")
                    if st.button("💾 Renommer Chantier", use_container_width=True):
                        n_cl = ch_new_n.strip()
                        if n_cl and n_cl != ch_a_renom:
                            conn = get_db_connection()
                            c = conn.cursor()
                            try:
                                c.execute("UPDATE chantiers_ref SET nom = ? WHERE nom = ?", (n_cl, ch_a_renom))
                                c.execute("UPDATE workers SET chantier_fixe = ? WHERE chantier_fixe = ?", (n_cl, ch_a_renom))
                                c.execute("UPDATE conducteur_chantiers SET chantier = ? WHERE chantier = ?", (n_cl, ch_a_renom))
                                c.execute("UPDATE pointages SET chantier = ? WHERE chantier = ?", (n_cl, ch_a_renom))
                                conn.commit()
                                conn.close()
                                st.session_state["sync_notif"] = f"✅ Chantier renommé vers {n_cl} !"
                                st.rerun()
                            except sqlite3.IntegrityError:
                                conn.close()
                                st.error("Ce nom existe déjà.")

                with tab_tch:
                    liste_tch = get_all_taches()
                    nouvelle_t = st.text_input("Nouvelle tâche :", key="in_add_tch_cfg")
                    if st.button("➕ Ajouter Tâche", type="primary", use_container_width=True):
                        n_t = nouvelle_t.strip()
                        if n_t:
                            conn = get_db_connection()
                            c = conn.cursor()
                            try:
                                c.execute("INSERT INTO taches_ref (nom) VALUES (?)", (n_t,))
                                conn.commit()
                                conn.close()
                                st.session_state["sync_notif"] = f"✅ Tâche {n_t} ajoutée !"
                                st.rerun()
                            except sqlite3.IntegrityError:
                                conn.close()
                                st.error("Cette tâche existe déjà.")
